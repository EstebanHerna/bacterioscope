"""Extract UZH/Dryad's source-reported resistance-mechanism labels.

Labels are preserved as dataset references. This module does not infer a
mechanism from AST categories and does not convert EUCAST interpretations to
CLSI interpretations.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

UZH_SOURCE_URL = "https://doi.org/10.5061/dryad.5dv41nsfj"
UZH_LICENSE = "CC0 1.0 (as declared by Dryad)"
UZH_STANDARD = "EUCAST"
UZH_STANDARD_VERSION = "Unreconciled: dataset README says 2022; repository records also say 2023"

_LABELS = {
    "none": "NO_MECHANISM_REPORTED",
    "no resistance mechanism": "NO_MECHANISM_REPORTED",
    "no resistance mechanism detected": "NO_MECHANISM_REPORTED",
    "no mechanism": "NO_MECHANISM_REPORTED",
    "esbl": "ESBL",
    "extended spectrum beta lactamase": "ESBL",
    "ampc": "AMPC",
    "ampc beta lactamase": "AMPC",
    "plasmid mediated ampc beta lactamase": "AMPC",
    "carbapenemase": "CARBAPENEMASE",
    "combination": "COMBINATION_UNSPECIFIED",
}

_OVERVIEW_FIELDS = {
    "image": ("Image", "image_id"),
    "organism": ("SPECIES", "organism", "species"),
    "none": ("None (Y/N)", "None"),
    "esbl": ("ESBL (Y/N)", "ESBL"),
    "ampc": ("AMPC (Y/N)", "AMPC"),
    "carbapenemase": ("CARBAPENEMASE (Y/N)", "CARBAPENEMASE"),
}


@dataclass(frozen=True)
class _SourceTable:
    headers: list[str]
    rows: list[dict[str, str]]
    source_asset: str
    source_author: str = ""
    representation: str = "csv"


def normalize_uzh_mechanism_label(raw_label: str) -> str | None:
    """Map only documented label aliases; return None for blank values."""

    cleaned = " ".join(re.sub(r"[^a-z0-9]+", " ", raw_label.casefold()).split())
    if not cleaned:
        return None
    return _LABELS.get(cleaned)


def _normalize_header(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", header.casefold())


def _find_column(headers: list[str], aliases: tuple[str, ...]) -> str | None:
    normalized = {_normalize_header(header): header for header in headers}
    for alias in aliases:
        if _normalize_header(alias) in normalized:
            return normalized[_normalize_header(alias)]
    return None


def _is_overview_workbook(headers: list[str]) -> bool:
    return all(_find_column(headers, aliases) is not None for aliases in _OVERVIEW_FIELDS.values())


def _cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _read_csv_stream(source: io.TextIOBase) -> tuple[list[str], list[dict[str, str]]]:
    sample = source.read(4096)
    source.seek(0)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(source, dialect=dialect)
    if not reader.fieldnames:
        raise ValueError("UZH CSV has no header")
    return list(reader.fieldnames), [dict(row) for row in reader]


def _read_xlsx_bytes(data: bytes, source_asset: str) -> _SourceTable:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Reading the UZH overview workbook requires openpyxl; install the project dev extra."
        ) from error

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        matching_sheets: list[tuple[list[str], list[dict[str, str]]]] = []
        for sheet in workbook.worksheets:
            row_iterator = sheet.iter_rows(values_only=True)
            header_row = next(row_iterator, None)
            if not header_row:
                continue
            headers = [_cell_text(value) for value in header_row]
            if not _is_overview_workbook(headers):
                continue

            rows: list[dict[str, str]] = []
            for row in row_iterator:
                if not any(value is not None for value in row):
                    continue
                rows.append(
                    {
                        header: _cell_text(row[index]) if index < len(row) else ""
                        for index, header in enumerate(headers)
                        if header
                    }
                )
            matching_sheets.append((headers, rows))

        if len(matching_sheets) != 1:
            if not matching_sheets:
                raise ValueError(
                    f"{source_asset} has no sheet with the expected Image, SPECIES, "
                    "None, ESBL, AMPC, and CARBAPENEMASE columns"
                )
            raise ValueError(f"{source_asset} has multiple UZH overview sheets")

        headers, rows = matching_sheets[0]
        return _SourceTable(
            headers=headers,
            rows=rows,
            source_asset=source_asset,
            source_author=_cell_text(workbook.properties.creator),
            representation="overview_workbook",
        )
    finally:
        workbook.close()


def _read_xlsx(input_path: Path) -> _SourceTable:
    if input_path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("UZH overview workbook exceeds the 20 MiB processing limit")
    return _read_xlsx_bytes(input_path.read_bytes(), input_path.name)


def _read_rows(input_path: Path) -> _SourceTable:
    if input_path.suffix.casefold() == ".xlsx":
        return _read_xlsx(input_path)

    if input_path.suffix.casefold() != ".zip":
        with input_path.open("r", encoding="utf-8-sig", newline="") as source:
            headers, rows = _read_csv_stream(source)
        return _SourceTable(headers=headers, rows=rows, source_asset=input_path.name)

    with ZipFile(input_path) as archive:
        files = [item for item in archive.infolist() if not item.is_dir()]
        csv_candidates = [
            item for item in files
            if not item.filename.startswith("__MACOSX/")
            and item.filename.casefold().endswith(".csv")
        ]
        preferred = [
            item for item in csv_candidates
            if Path(item.filename).name.casefold() == "measurements.csv"
        ]
        if preferred:
            csv_candidates = preferred
        if len(csv_candidates) == 1:
            csv_info = csv_candidates[0]
            if csv_info.file_size > 20 * 1024 * 1024:
                raise ValueError("UZH measurements CSV exceeds the 20 MiB processing limit")
            with archive.open(csv_info) as binary_source:
                with io.TextIOWrapper(binary_source, encoding="utf-8-sig", newline="") as source:
                    headers, rows = _read_csv_stream(source)
            return _SourceTable(
                headers=headers,
                rows=rows,
                source_asset=csv_info.filename,
            )
        if csv_candidates:
            raise ValueError("UZH archive must contain one unambiguous measurements CSV")

        workbook_candidates = [
            item for item in files
            if not item.filename.startswith("__MACOSX/")
            and item.filename.casefold().endswith(".xlsx")
        ]
        matching_workbooks: list[_SourceTable] = []
        for workbook_info in workbook_candidates:
            if workbook_info.file_size > 20 * 1024 * 1024:
                continue
            try:
                workbook = _read_xlsx_bytes(archive.read(workbook_info), workbook_info.filename)
            except ValueError:
                continue
            if _is_overview_workbook(workbook.headers):
                matching_workbooks.append(workbook)

        if len(matching_workbooks) == 1:
            return matching_workbooks[0]
        if len(matching_workbooks) > 1:
            raise ValueError("UZH archive contains multiple phenotype overview workbooks")

        xlsx_names = [item.filename for item in workbook_candidates[:5]]
        raise ValueError(
            "UZH archive has no measurements CSV and no workbook with the expected "
            f"phenotype columns; workbooks found: {xlsx_names}"
        )


def extract_uzh_phenotype_labels(
    csv_path: Path | str,
    output_path: Path | str | None = None,
) -> list[dict[str, str]]:
    """Extract source-reported UZH labels and optionally write a CSV.

    ``csv_path`` may be Dryad's ``Tables.zip``, its documented measurements
    CSV, or the phenotype overview workbook found in the current archive.
    Repeated measurements are collapsed by isolate. Raw workbook indicators
    and unknown statuses remain visible. Workbook summary labels are not
    assigned reference status because their relationship to the paper's
    routine-diagnostic labels is unverified. This function never infers a
    missing mechanism or converts EUCAST labels to CLSI.
    """

    input_path = Path(csv_path)
    source_table = _read_rows(input_path)
    source_sha256 = hashlib.sha256(input_path.read_bytes()).hexdigest()
    headers = source_table.headers
    rows = source_table.rows
    records: list[dict[str, str]] = []

    if source_table.representation == "overview_workbook":
        image_column = _find_column(headers, _OVERVIEW_FIELDS["image"])
        organism_column = _find_column(headers, _OVERVIEW_FIELDS["organism"])
        none_column = _find_column(headers, _OVERVIEW_FIELDS["none"])
        mechanism_columns = {
            "ESBL": _find_column(headers, _OVERVIEW_FIELDS["esbl"]),
            "AMPC": _find_column(headers, _OVERVIEW_FIELDS["ampc"]),
            "CARBAPENEMASE": _find_column(headers, _OVERVIEW_FIELDS["carbapenemase"]),
        }
        if not image_column or not organism_column or not none_column or any(
            column is None for column in mechanism_columns.values()
        ):
            raise ValueError("UZH overview workbook is missing a required phenotype column")

        seen_image_ids: set[str] = set()
        seen_image_join_keys: set[str] = set()
        for row in rows:
            image_id = (row.get(image_column) or "").strip()
            if not image_id:
                raise ValueError("UZH overview workbook contains a row without an Image ID")
            if image_id.casefold() in seen_image_ids:
                raise ValueError(f"duplicate UZH Image ID: {image_id}")
            seen_image_ids.add(image_id.casefold())
            image_join_key = image_id.rstrip(".").strip()
            if image_join_key.casefold() in seen_image_join_keys:
                raise ValueError(
                    f"UZH Image IDs collide after trailing-period normalization: {image_id}"
                )
            seen_image_join_keys.add(image_join_key.casefold())

            raw_indicators: list[str] = []
            unmapped: list[str] = []
            mapped: list[str] = []
            none_value = (row.get(none_column) or "").strip()
            raw_indicators.append(f"{none_column}={none_value}")
            if none_value.casefold() == "yes":
                mapped.append("NO_MECHANISM_REPORTED")
            elif none_value.casefold() != "no":
                unmapped.append(f"{none_column}={none_value or '<blank>'}")

            positive_mechanisms: list[str] = []
            for code, column in mechanism_columns.items():
                if column is None:
                    raise ValueError("UZH overview workbook is missing a mechanism column")
                value = (row.get(column) or "").strip()
                raw_indicators.append(f"{column}={value}")
                if value.casefold() == "yes":
                    positive_mechanisms.append(code)
                    mapped.append(code)
                elif value.casefold() != "no":
                    unmapped.append(f"{column}={value or '<blank>'}")

            conflict = none_value.casefold() == "yes" and bool(positive_mechanisms)
            if conflict:
                label_status = "conflict"
                unmapped.append("None indicator conflicts with positive mechanism indicator")
            elif unmapped:
                label_status = "partial"
            else:
                label_status = "complete"

            records.append(
                {
                    "source": "dryad_uzh_sirscan",
                    "sample_id": image_id,
                    "source_identifier_type": "source_image_id",
                    "image_id": image_id,
                    "image_join_key": image_join_key,
                    "organism": (row.get(organism_column) or "").strip(),
                    "source_asset": source_table.source_asset,
                    "source_author": source_table.source_author,
                    "raw_labels": json.dumps(raw_indicators, ensure_ascii=False),
                    "source_reported_codes": json.dumps(mapped),
                    "reference_label_codes": json.dumps([]),
                    "unmapped_labels": json.dumps(unmapped, ensure_ascii=False),
                    "label_status": label_status,
                    "label_provenance_status": "unverified_workbook_summary",
                    "evidence_type": (
                        "unverified_workbook_summary; possible model-output; "
                        "not ground truth until source role is confirmed"
                    ),
                    "standard": UZH_STANDARD,
                    "standard_version_note": UZH_STANDARD_VERSION,
                    "source_url": UZH_SOURCE_URL,
                    "license": UZH_LICENSE,
                    "source_file_sha256": source_sha256,
                }
            )

        records.sort(key=lambda record: record["sample_id"].casefold())
    else:
        isolate_column = _find_column(
            headers, ("SampleID", "sample_id", "isolate_id", "isolate")
        )
        mechanism_column = _find_column(
            headers,
            ("ResistanceMechanism", "mechanism", "phenotype", "resistance_phenotype"),
        )
        if isolate_column is None or mechanism_column is None:
            raise ValueError(
                "UZH table must contain an isolate identifier and a resistance-mechanism label"
            )
        organism_column = _find_column(headers, ("Species", "SPECIES", "Organism"))

        grouped: dict[str, set[str]] = defaultdict(set)
        organism_by_isolate: dict[str, set[str]] = defaultdict(set)
        for row in rows:
            isolate_id = (row.get(isolate_column) or "").strip()
            raw_label = (row.get(mechanism_column) or "").strip()
            if not isolate_id or not raw_label:
                continue
            grouped[isolate_id].add(raw_label)
            if organism_column and (organism := (row.get(organism_column) or "").strip()):
                organism_by_isolate[isolate_id].add(organism)

        for isolate_id in sorted(grouped):
            raw_labels = sorted(grouped[isolate_id], key=str.casefold)
            mapped_codes = {
                normalized
                for label in raw_labels
                if (normalized := normalize_uzh_mechanism_label(label)) is not None
            }
            mapped = sorted(mapped_codes)
            unmapped = sorted(
                {label for label in raw_labels if normalize_uzh_mechanism_label(label) is None},
                key=str.casefold,
            )
            organisms = sorted(organism_by_isolate[isolate_id], key=str.casefold)
            records.append(
                {
                    "source": "dryad_uzh_sirscan",
                    "sample_id": isolate_id,
                    "source_identifier_type": "sample_id",
                    "image_id": "",
                    "image_join_key": "",
                    "organism": organisms[0] if len(organisms) == 1 else "",
                    "source_asset": source_table.source_asset,
                    "source_author": source_table.source_author,
                    "raw_labels": json.dumps(raw_labels, ensure_ascii=False),
                    "source_reported_codes": json.dumps(mapped),
                    "reference_label_codes": json.dumps(mapped),
                    "unmapped_labels": json.dumps(unmapped, ensure_ascii=False),
                    "label_status": "partial" if unmapped else "complete",
                    "label_provenance_status": "documented_measurement_table",
                    "evidence_type": (
                        "source_reported_mechanism_label; not inferred by BacterioScope"
                    ),
                    "standard": UZH_STANDARD,
                    "standard_version_note": UZH_STANDARD_VERSION,
                    "source_url": UZH_SOURCE_URL,
                    "license": UZH_LICENSE,
                    "source_file_sha256": source_sha256,
                }
            )

    if output_path is not None:
        destination = Path(output_path)
        destination.parent.mkdir(parents=True, exist_ok=True)
        fields = [
            "source",
            "sample_id",
            "source_identifier_type",
            "image_id",
            "image_join_key",
            "organism",
            "source_asset",
            "source_author",
            "raw_labels",
            "source_reported_codes",
            "reference_label_codes",
            "unmapped_labels",
            "label_status",
            "label_provenance_status",
            "evidence_type",
            "standard",
            "standard_version_note",
            "source_url",
            "license",
            "source_file_sha256",
        ]
        with destination.open("w", encoding="utf-8", newline="") as output:
            writer = csv.DictWriter(output, fieldnames=fields)
            writer.writeheader()
            writer.writerows(records)
    return records
